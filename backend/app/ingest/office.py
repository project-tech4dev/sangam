"""Text from Word, Excel and text files, for an upload's content_md and the readers. Pure: bytes in, text out."""

from __future__ import annotations

import io

SHEET_ROWS = 200  # an upload's text keeps this many rows a sheet; the vault reader pages the rest
LINE_MAX = 2_000  # a text file's line longer than this is split, so one minified line cannot fill a read


def docx_paragraphs(data: bytes) -> list[str]:
    """Each paragraph, and each table row as `a | b | c`, in document order; empty ones dropped."""
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(io.BytesIO(data))
    out: list[str] = []
    for block in doc.element.body.iterchildren():
        tag = block.tag.rsplit("}", 1)[-1]
        if tag == "p":
            text = Paragraph(block, doc).text.strip()
            if text:
                out.append(text)
        elif tag == "tbl":
            for row in Table(block, doc).rows:
                # A merged cell repeats across the row as the same element; keep it once.
                cells = list({id(c._tc): c.text.strip() for c in row.cells}.values())
                if any(cells):
                    out.append(" | ".join(cells))
    return out


def docx_range(data: bytes, start: int = 1, end: int | None = None) -> tuple[str, int]:
    """(paragraphs start..end, each led by [para N]; how many the file has), so a Word file pages like a PDF."""
    paras = docx_paragraphs(data)
    last = min(end or len(paras), len(paras))
    return "\n".join(f"[para {n}] {paras[n - 1]}" for n in range(max(start, 1), last + 1)), len(paras)


def text_lines(data: bytes) -> list[str]:
    """A markdown or text file's lines, a very long one split at a space, so every line fits one read."""
    lines: list[str] = []
    for raw in data.decode("utf-8", errors="replace").splitlines():
        line = raw.rstrip()
        while len(line) > LINE_MAX:
            cut = line.rfind(" ", 0, LINE_MAX)
            cut = cut if cut > 0 else LINE_MAX
            lines.append(line[:cut])
            line = line[cut:].lstrip()
        lines.append(line)
    return lines


def text_range(data: bytes, start: int = 1, end: int | None = None) -> tuple[str, int]:
    """(lines start..end, each led by [line N], blank ones skipped; how many the file has), paged like a PDF."""
    lines = text_lines(data)
    last = min(end or len(lines), len(lines))
    return "\n".join(f"[line {n}] {lines[n - 1]}" for n in range(max(start, 1), last + 1)
                     if lines[n - 1].strip()), len(lines)


def xlsx_markdown(data: bytes) -> str:
    """Every sheet as a heading and its first rows, cells joined by `|`."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    lines: list[str] = []
    try:
        for ws in wb.worksheets:
            lines.append(f"## {ws.title}")
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= SHEET_ROWS:
                    lines.append(f"… first {SHEET_ROWS} rows only")
                    break
                cells = ["" if v is None else str(v).strip() for v in row]
                if any(cells):
                    lines.append(" | ".join(cells))
    finally:
        wb.close()
    return "\n".join(lines)
