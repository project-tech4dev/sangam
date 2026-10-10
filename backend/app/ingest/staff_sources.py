"""Staff-only source workbooks: where they live, how they are read, how a fact from one is cited.

Confidential: nothing here logs or returns an error carrying a cell value."""

from __future__ import annotations

import asyncio
import hashlib
import io
import logging
import re
from collections import OrderedDict
from collections.abc import Iterable
from pathlib import PurePath
from typing import Any

from app.ingest.sources import SOURCES
from app.integrations import storage
from app.repositories import documents as documents_repo
from app.repositories import funders as funders_repo

logger = logging.getLogger(__name__)

DOC_TYPE = "source_dataset"
SOURCE_URL_PREFIX = "https://staff-sources.sangam.invalid/"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def dataset_id(filename: str) -> str:
    """A file's identity is its base name, so a corrected re-upload replaces it."""
    return hashlib.sha256(PurePath(filename).name.encode()).hexdigest()[:12]


def object_key(source_key: str, filename: str) -> str:
    return f"{source_key}/{dataset_id(filename)}.xlsx"


def source_url(filename: str) -> str:
    """http(s) as findings require; resolves nowhere (RFC 2606) and names no file."""
    return SOURCE_URL_PREFIX + dataset_id(filename)


def is_staff_source_url(url: str | None) -> bool:
    return bool(url) and str(url).startswith(SOURCE_URL_PREFIX)


def title(source_key: str, filename: str) -> str:
    return f"Staff dataset · {source_key} · {dataset_id(filename)}"


_ID = re.compile(r"^[0-9a-f]{12}$")
PICKED_VERSION = "select_source_files/1"


_BUCKET = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._-]{0,62}$")


def picked_id(path: str) -> str:
    """A picked object's id: the script's own when it wrote the key, so one object keeps one source_url."""
    stem = PurePath(path).stem
    if _ID.match(stem) and path.startswith(storage.STAFF_SOURCES + "/"):
        return stem
    return hashlib.sha256(path.encode()).hexdigest()[:12]


def split_any(path: str) -> tuple[str, str] | None:
    """(bucket, key) for a path in any bucket; storage.split stays limited to the buckets we hand URLs for."""
    bucket, _, key = (path or "").partition("/")
    if not _BUCKET.match(bucket) or not key or any(seg in ("", ".", "..") for seg in key.split("/")):
        return None
    return bucket, key


def check_pick(path: str) -> tuple[str, str, str] | str:
    """(bucket, key, source name) for a selectable path, or why it is not one."""
    parts = split_any(path)
    if parts is None:
        return f"{path!r}: not a <bucket>/<file> path from list_bucket_files."
    bucket, key = parts
    # Only the staff bucket: a Doc Vault workbook is an org's private file, not a staff dataset.
    if bucket != storage.STAFF_SOURCES:
        return f"{path!r}: only files in the {storage.STAFF_SOURCES} bucket can be read as staff data."
    if not key.lower().endswith(".xlsx"):
        return f"{path!r}: only .xlsx workbooks can be read."
    # A registry-key folder names the source (and so its unit); any other file is named by its bucket.
    folder = key.split("/", 1)[0] if "/" in key else ""
    return bucket, key, folder if folder in SOURCES else bucket


async def register(pool: Any, path: str) -> int | str:
    """Download a picked workbook, write its documents row, and keep the bytes for the readers."""
    checked = check_pick(path)
    if isinstance(checked, str):
        return checked
    bucket, key, source_name = checked
    data = await storage.download(bucket, key, max_bytes=storage.STAFF_SOURCES_MAX_BYTES)
    if data is None:
        return f"{path!r}: could not be downloaded."
    try:
        await asyncio.to_thread(_open_check, data)
    except Exception as exc:  # noqa: BLE001 - type only: a message could quote a cell
        return f"{path!r}: not a readable .xlsx workbook ({type(exc).__name__})."
    digest = storage.content_hash(data)
    qualified = storage.qualify(bucket, key)
    ident = picked_id(qualified)
    document_id, _ = await documents_repo.upsert_source_dataset(
        pool, title=f"Staff dataset · {source_name} · {ident}", source_key=source_name,
        storage_path=qualified, source_url=SOURCE_URL_PREFIX + ident,
        content_hash=digest, extractor_version=PICKED_VERSION,
    )
    _remember((qualified, digest), data)
    return document_id


def _open_check(data: bytes) -> None:
    from openpyxl import load_workbook

    load_workbook(io.BytesIO(data), read_only=True).close()


HEADER_ROWS = 3            # top rows of a sheet shown with any match
SECTION_ROWS = 40          # rows a matching numbered heading brings with it
MAX_CHARS = 60_000         # ~15k tokens of tool output
# Matched as whole words after folding, so "tel." is "tel" and "e-mail" is "e mail".
CONTACT_HEADERS = ("phone", "mobile", "email", "e mail", "contact person", "address", "telephone", "tel",
                   "contact no", "contact number", "landline", "fax", "whatsapp")
_SUFFIXES = {"ltd", "limited", "pvt", "private"}
_NON_WORD = re.compile(r"[^\w]+")
_NUMBERED = re.compile(r"^\s*\d+\s*[.)]\s*\S")
_EMAIL = re.compile(r"\S+@\S+\.\w+")
_PHONE = re.compile(r"\+?\d[\d \-()]{8,}\d")
_YEARS = re.compile(r"^(?:(?:19|20)\d\d[\s\-–/]*)+$")
_CACHE_MAX_BYTES = 200 * 1024 * 1024
_cache: OrderedDict[tuple[str, str], bytes] = OrderedDict()
_locks: dict[str, asyncio.Lock] = {}


def _words(value: Any) -> str:
    return " ".join(_NON_WORD.sub(" ", str(value).casefold()).split())


def search_terms(names: Iterable[str | None]) -> list[str]:
    terms: list[str] = []
    for name in names:
        words = _words(name or "").split()
        while words and words[-1] in _SUFFIXES:
            words.pop()
        if words and words[0] == "the":
            words.pop(0)
        term = " ".join(words)
        if len(term) >= 3 and term not in terms:
            terms.append(term)
    return terms


def _hit(value: Any, terms: list[str]) -> bool:
    if not isinstance(value, str):
        return False
    padded = f" {_words(value)} "
    return any(f" {t} " in padded for t in terms)


def _contact_columns(values: tuple) -> set[int]:
    return {c for c, v in enumerate(values, start=1)
            if isinstance(v, str) and any(f" {h} " in f" {_words(v)} " for h in CONTACT_HEADERS)}


def _block_starts(top: tuple, second: tuple) -> list[int]:
    """Columns that open a side-by-side block: a sparse title row (one name per block) over a dense header row."""
    top_cols = [c for c, v in enumerate(top, start=1) if v not in (None, "")]
    second_filled = sum(1 for v in second if v not in (None, ""))
    contiguous = top_cols == list(range(top_cols[0], top_cols[0] + len(top_cols))) if top_cols else True
    if len(top_cols) < 2 or contiguous or second_filled < 2 * len(top_cols):
        return []
    return top_cols


def _in_blocks(values: tuple, terms: list[str], starts: list[int]) -> set[int] | None:
    """The columns of the blocks a row's matches sit in, so a row never shows another block's cells."""
    owners = set()
    for c, v in enumerate(values, start=1):
        if _hit(v, terms):
            owners.add(max((s for s in starts if s <= c), default=starts[0]))
    if not owners:
        return None
    bounds = [*starts[1:], len(values) + 1]
    return {c for s, end in zip(starts, bounds) if s in owners for c in range(s, end)}


def scan(data: bytes, terms: list[str], only_sheet: str = "") -> list[tuple[str, list[tuple[int, tuple, set[int] | None]], set[int]]]:
    """Per sheet with a match: its title, the rows to show with the columns each may show, and contact columns."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = []
    try:
        for ws in wb.worksheets:
            if only_sheet and only_sheet.casefold() not in ws.title.casefold():
                continue
            ws.reset_dimensions()  # some exports carry a wrong <dimension>; read_only trusts it otherwise
            rows: dict[int, tuple[tuple, set[int] | None]] = {}
            drop: set[int] = set()
            top: tuple = ()
            starts: list[int] = []
            numbered = heading = under = None
            section_left, matched, seen = 0, False, 0
            for r, values in enumerate(ws.iter_rows(values_only=True), start=1):
                filled = [v for v in values if v not in (None, "")]
                if not filled:
                    continue
                seen += 1
                hit = any(_hit(v, terms) for v in filled)
                if seen == 1:
                    top = values
                elif seen == 2:
                    starts = _block_starts(top, values)
                # Shown always; a match here still counts (data can start on row 2). A block sheet's data starts on row 3.
                if seen <= (2 if starts else HEADER_ROWS):
                    rows[r] = (values, None)
                    drop |= _contact_columns(values)
                    matched = matched or hit
                    continue
                if heading and under is None and r == heading[0] + 1:
                    under = (r, values)
                    drop |= _contact_columns(values)
                if len(filled) == 1:
                    heading, under = (r, values), None
                    if _NUMBERED.match(str(filled[0])):
                        numbered, section_left = (r, values), (SECTION_ROWS if hit else 0)
                if hit or section_left:
                    for ctx in (numbered, heading, under):
                        if ctx:
                            rows.setdefault(ctx[0], (ctx[1], None))
                    rows[r] = (values, _in_blocks(values, terms, starts) if starts and hit else None)
                    matched = matched or hit
                    section_left = max(section_left - 1, 0)
            if matched:
                out.append((ws.title, [(r, v, keep) for r, (v, keep) in sorted(rows.items())], drop))
    finally:
        wb.close()
    return out


def _phone(text: str) -> bool:
    """A run of 10+ digits in a text cell (numbers are amounts); a run of years is not one."""
    return any(sum(ch.isdigit() for ch in m) >= 10 and not _YEARS.match(m) for m in _PHONE.findall(text))


def _cell(value: Any) -> str:
    text = " ".join(str(value).split())
    if _EMAIL.search(text):
        return "[email removed]"
    if isinstance(value, str) and _phone(text):
        return "[phone removed]"
    return text[:300]


def render(doc: dict[str, Any], sheets: list[tuple[str, list[tuple[int, tuple, set[int] | None]], set[int]]]) -> str:
    from openpyxl.utils import get_column_letter

    lines = [f"## document {doc['id']} · {doc['source_name']}",
             f"A finding from it: source_url {doc['source_url']} · source_name {doc['source_name']}"]
    for title, rows, drop in sheets:
        lines.append(f"### sheet {title!r}")
        for r, values, keep in rows:
            cells = [(f"{get_column_letter(c)}{r}", v) for c, v in enumerate(values, start=1)
                     if v not in (None, "") and c not in drop and (keep is None or c in keep)]
            if cells:
                body = " · ".join(f"{ref}={_cell(v)!r}" for ref, v in cells)
                lines.append(f"[source:{doc['id']}#{title}!{cells[0][0]}] {body}")
    return "\n".join(lines)


async def _workbook(doc: dict[str, Any]) -> bytes | None:
    """The workbook's bytes, downloaded once per process and per upload (content_hash)."""
    parts = split_any(doc["storage_path"] or "")
    if parts is None:
        return None
    key = (doc["storage_path"], doc["content_hash"] or "")
    async with _locks.setdefault(doc["storage_path"], asyncio.Lock()):
        if key in _cache:
            _cache.move_to_end(key)
            return _cache[key]
        data = await storage.download(*parts, max_bytes=storage.STAFF_SOURCES_MAX_BYTES)
        if data is None:
            return None
        _remember(key, data)
        return data


def _remember(key: tuple[str, str], data: bytes) -> None:
    _cache[key] = data
    _cache.move_to_end(key)
    while len(_cache) > 1 and sum(len(b) for b in _cache.values()) > _CACHE_MAX_BYTES:
        _cache.popitem(last=False)


async def read_rows(pool: Any, name: str, *, sheet: str = "", documents: list[int] | None = None) -> str:
    name = (name or "").strip()
    asked = search_terms([name])
    funders = await funders_repo.name_variants(pool, match=asked[0]) if asked else []
    terms = search_terms([name, *(n for f in funders for n in (f["name"], *(f["aliases"] or [])))])
    if not terms:
        return "read_source_rows needs a name of at least three letters."
    datasets = await documents_repo.list_source_datasets(pool, ids=documents or None)
    if not datasets:
        if documents:
            return f"None of documents {documents} is a staff dataset. Nothing was read."
        return "No staff dataset is registered: pick files with select_source_files. Nothing was read."
    parts: list[str] = []
    problems: list[str] = []
    for doc in datasets:
        try:
            data = await _workbook(doc)
        except Exception as exc:  # noqa: BLE001 - one file failing is listed, not fatal
            data, kind = None, type(exc).__name__
        else:
            kind = ""
        if data is None:
            if kind:
                problems.append(f"document {doc['id']} ({doc['source_name']}) could not be read ({kind}).")
                continue
            problems.append(f"document {doc['id']} ({doc['source_name']}) could not be downloaded.")
            continue
        try:
            sheets = await asyncio.to_thread(scan, data, terms, sheet)
        except Exception as exc:  # noqa: BLE001 - type only: a message could quote a cell
            problems.append(f"document {doc['id']} ({doc['source_name']}) could not be read ({type(exc).__name__}).")
            continue
        if sheets:
            parts.append(render(doc, sheets))
    logger.info("read_source_rows: %d of %d datasets matched", len(parts), len(datasets))
    if len(problems) == len(datasets):
        return "\n".join([f"read_source_rows could not read {len(datasets)} staff file(s). Nothing was searched.",
                          *problems])
    head = [f"read_source_rows: {len(parts)} of {len(datasets)} staff datasets name one of: {'; '.join(terms)}."]
    if not parts:
        head.append("No row matched. The filer is often the parent company, not the foundation: try its "
                    "registered name.")
        return "\n".join(head + problems)
    head.append("Cite a fact by the marker of the cell that holds it: [source:<document>#<sheet>!<cell>]. "
                "Amounts are as printed; the sheet's unit is in \"The staff sheets\" in your instructions.")
    text = "\n\n".join(["\n".join(head), *parts])
    tail = "\n\n" + "\n".join(problems) if problems else ""
    room = MAX_CHARS - len(tail)
    if len(text) > room:  # cut rows, never the failure lines
        cut = text.rfind("\n", 0, room)
        text = (text[:cut] + f"\n… output cut at {MAX_CHARS} characters. Call again with `sheet` "
                "(e.g. one fiscal year) to read the rest.")
    return text + tail



BROWSE_MAX_ROWS = 200      # rows one read_workbook page may return


def _browse(data: bytes, only_sheet: str, start_row: int, limit: int) -> tuple[list, list[str]]:
    """Every sheet's header rows and size, or one sheet's header rows plus rows start_row.. of it."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheets, notes = [], []
    try:
        for ws in wb.worksheets:
            if only_sheet and only_sheet.casefold() not in ws.title.casefold():
                continue
            ws.reset_dimensions()
            rows: list[tuple[int, tuple, None]] = []
            drop: set[int] = set()
            seen = width = 0
            more = None
            for r, values in enumerate(ws.iter_rows(values_only=True), start=1):
                if not any(v not in (None, "") for v in values):
                    continue
                seen += 1
                width = max(width, max(c for c, v in enumerate(values, start=1) if v not in (None, "")))
                if seen <= HEADER_ROWS:
                    drop |= _contact_columns(values)
                    rows.append((r, values, None))
                elif only_sheet and start_row <= r < start_row + limit:
                    rows.append((r, values, None))
                elif only_sheet and r >= start_row + limit and more is None:
                    more = r
            sheets.append((ws.title, rows, drop))
            notes.append(f"sheet {ws.title!r}: {seen} filled rows × {width} columns"
                         + (f"; more rows: call again with start_row={more}" if more else ""))
            if only_sheet:
                break
    finally:
        wb.close()
    return sheets, notes


async def browse(pool: Any, documents: list[int], *, sheet: str = "", start_row: int = 1, limit: int = 100) -> str:
    """read_workbook: an outline of each picked workbook, or one page of one sheet's rows."""
    if not documents:
        return "read_workbook needs the document ids select_source_files returned. Nothing was read."
    datasets = await documents_repo.list_source_datasets(pool, ids=documents)
    if not datasets:
        return f"None of documents {documents} is a staff dataset. Nothing was read."
    start_row, limit = max(1, start_row), max(1, min(limit, BROWSE_MAX_ROWS))
    parts: list[str] = []
    problems: list[str] = []
    for doc in datasets:
        try:
            data = await _workbook(doc)
            if data is None:
                problems.append(f"document {doc['id']} ({doc['source_name']}) could not be downloaded.")
                continue
            sheets, notes = await asyncio.to_thread(_browse, data, sheet, start_row, limit)
        except Exception as exc:  # noqa: BLE001 - type only: a message could quote a cell
            problems.append(f"document {doc['id']} ({doc['source_name']}) could not be read ({type(exc).__name__}).")
            continue
        if not sheets:
            problems.append(f"document {doc['id']} has no sheet matching {sheet!r}.")
            continue
        parts.append("\n".join([render(doc, sheets), *notes]))
    head = ("Outline: each sheet's header rows and size. Call again with `sheet` to page through its rows."
            if not sheet else f"Rows {start_row}–{start_row + limit - 1} of sheet {sheet!r}, under its header rows.")
    head += " Contact columns are removed. Cite a cell by its [source:<document>#<sheet>!<cell>] marker."
    text = "\n\n".join([head, *parts])
    tail = "\n\n" + "\n".join(problems) if problems else ""
    room = MAX_CHARS - len(tail)
    if len(text) > room:  # cut rows, never the failure lines
        cut = text.rfind("\n", 0, room)
        text = text[:cut] + f"\n… output cut at {MAX_CHARS} characters. Call again with a smaller `limit`."
    return text + tail

_MARKER = re.compile(r"\[source:(\d+)#[^\]]+![A-Z]{1,3}\d+\]", re.IGNORECASE)


def hold_unmarked(
    findings: list[dict[str, Any]], dataset_ids: dict[str, int]
) -> tuple[list[dict[str, Any]], list[str]]:
    """Promotable findings, and a held reason for each staff finding not marked with its own dataset's cell."""
    kept: list[dict[str, Any]] = []
    held: list[str] = []
    for f in findings:
        url = f.get("source_url")
        if not is_staff_source_url(url):
            kept.append(f)
            continue
        ids = {int(m) for m in _MARKER.findall(str(f.get("evidence", "")))}
        if not ids:
            held.append(f"{f.get('concept')}: cites a staff dataset without its [source:…] marker — quote "
                        f"the row's line with the marker of the cell that holds the fact")
        elif url not in dataset_ids:
            held.append(f"{f.get('concept')}: its source_url names no staff dataset on file")
        elif dataset_ids[url] not in ids:
            held.append(f"{f.get('concept')}: the citation names a different staff dataset")
        else:
            kept.append(f)
    return kept, held


# Checking a staff figure against its cell. A marker proves which cell was cited, not that the number
# was read off it right: a misread value or a wrong unit would otherwise be promoted as a public fact.

_CELL = re.compile(r"\[source:(\d+)#(.+?)!([A-Z]{1,3})(\d+)\]", re.IGNORECASE)
_NUMBER = re.compile(r"\d[\d,]*(?:\.\d+)?")
_UNIT_WORDS = (("lakh", "lakh"), ("lac", "lakh"), ("crore", "crore"), ("cr", "crore"))
# A sheet whose plain numbers are always in one unit. The CSO workbook's units vary by sheet and cell.
SOURCE_UNITS = {"mca_csr": "crore"}


def cell_amount(value: Any) -> tuple[float, str | None] | None:
    """A cell as (number as printed, the unit it states or None); None when it holds no number."""
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value), None
    text = str(value)
    found = _NUMBER.search(text)
    if not found:
        return None
    words = _words(text).split()
    unit = next((u for w, u in _UNIT_WORDS for word in words if word == w or word.startswith(w + "s")), None)
    if unit is None and re.search(r"\dcr\b", text, re.IGNORECASE):
        unit = "crore"
    return float(found.group().replace(",", "")), unit


def _same(a: float, b: float) -> bool:
    return abs(a - b) <= max(0.005, 0.0005 * abs(b))


def _money_fields(finding: dict[str, Any]) -> dict[str, dict[str, Any]]:
    return {k: v for k, v in finding.items()
            if isinstance(v, dict) and {"value", "unit", "text"} <= v.keys()}


def _read_rows(data: bytes, wanted: dict[str, set[int]]) -> dict[tuple[str, int], tuple]:
    """The cited rows, one streaming pass per sheet. A sheet that is not in the workbook is simply absent."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out: dict[tuple[str, int], tuple] = {}
    try:
        for title, numbers in wanted.items():
            if title not in wb.sheetnames:
                continue
            ws = wb[title]
            ws.reset_dimensions()
            for r, values in enumerate(ws.iter_rows(min_row=min(numbers), max_row=max(numbers),
                                                    values_only=True), start=min(numbers)):
                if r in numbers:
                    out[(title, r)] = values
    finally:
        wb.close()
    return out


def figure_mismatches(finding: dict[str, Any], doc_id: int, source_key: str,
                      rows: dict[tuple[str, int], tuple]) -> list[str]:
    """Each amount must be printed in a row the finding cites, in the unit the cell or sheet states."""
    cited = [(sheet, int(row)) for d, sheet, _col, row in _CELL.findall(str(finding.get("evidence", "")))
             if int(d) == doc_id]
    cells = [cell_amount(v) for key in cited for v in rows.get(key, ())]
    cells = [c for c in cells if c is not None]
    problems = []
    for name, money in _money_fields(finding).items():
        value, unit = float(money["value"]), money["unit"]
        ok = False
        for number, stated in cells:
            if not _same(value, number):
                continue
            expected = stated or SOURCE_UNITS.get(source_key)
            if expected is None or expected == unit:
                ok = True
                break
        if not ok:
            refs = ", ".join(f"{s}!row {r}" for s, r in cited) or "no readable row"
            problems.append(f"{finding.get('concept')}: {name} is not printed, in that unit, in the cited row ({refs}) "
                            f"— re-read the row and save the amount as printed with its unit")
    return problems


async def hold_mismatched(
    findings: list[dict[str, Any]], datasets: dict[str, dict[str, Any]]
) -> tuple[list[dict[str, Any]], list[str]]:
    """Promotable findings, and a held reason for each staff amount its cited row does not bear out."""
    kept: list[dict[str, Any]] = []
    held: list[str] = []
    rows_by_doc: dict[int, dict[tuple[str, int], tuple] | None] = {}
    for f in findings:
        doc = datasets.get(f.get("source_url") or "")
        if doc is None or not is_staff_source_url(f.get("source_url")) or not _money_fields(f):
            kept.append(f)
            continue
        if doc["id"] not in rows_by_doc:
            wanted: dict[str, set[int]] = {}
            for g in findings:
                if g.get("source_url") == f.get("source_url"):
                    for d, sheet, _col, row in _CELL.findall(str(g.get("evidence", ""))):
                        if int(d) == doc["id"]:
                            wanted.setdefault(sheet, set()).add(int(row))
            try:
                data = await _workbook(doc)
                rows_by_doc[doc["id"]] = await asyncio.to_thread(_read_rows, data, wanted) if data else None
            except Exception as exc:  # noqa: BLE001 - type only: a message could quote a cell
                logger.warning("staff figure check could not open document %s (%s)", doc["id"], type(exc).__name__)
                rows_by_doc[doc["id"]] = None
        rows = rows_by_doc[doc["id"]]
        if rows is None:
            held.append(f"{f.get('concept')}: its staff workbook could not be re-opened to check the figure")
            continue
        problems = figure_mismatches(f, doc["id"], doc.get("source_name") or "", rows)
        if problems:
            held.extend(problems)
        else:
            kept.append(f)
    return kept, held
