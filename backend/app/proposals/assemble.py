"""The proposal as one document: its written parts in reading order, without source marks, and as a Word file. Pure."""

from __future__ import annotations

import io
import re
from typing import Any

from app.proposals import rubric, trace

DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_MARK = re.compile(r"[ \t]*⟦k\d+⟧")
_BULLET = re.compile(r"^\s*[-•*]\s+")
_LEFT = re.compile(r"(\[(?:Add|You fill):[^\]]*\])")  # what the person still has to write, highlighted


def plain(text: str | None) -> str:
    return _MARK.sub("", (text or "").replace("**", "")).strip()


def parts(work: dict[str, Any]) -> list[tuple[rubric.Section, str]]:
    """Each written part as it stands, the Summary first and the rest in the rubric's order."""
    held = work.get("sections") or {}
    written = [(s, plain(held[s.code]["text"])) for s in rubric.sections_for(bool(work.get("earlier_phase")),
                                                                              work.get("template"))
               if (held.get(s.code) or {}).get("text")]
    return sorted(written, key=lambda p: p[0].code != "A1")


def final(work: dict[str, Any]) -> list[tuple[rubric.Section, str]]:
    """Each part as it goes out: the smoothed wording, unless the part changed after it was smoothed."""
    polished = work.get("polished") or {}
    smooth, revs = polished.get("parts") or {}, polished.get("revs") or {}
    held = work.get("sections") or {}
    return [(s, smooth[s.code] if s.code in smooth and revs.get(s.code) == held[s.code].get("rev", 0) else text)
            for s, text in parts(work)]


def snapshot(work: dict[str, Any]) -> dict[str, dict[str, Any]]:
    """The facts the text cites, as they stood when a version was saved."""
    facts = work.get("facts") or {}
    cited = {k for s in (work.get("sections") or {}).values() for k in trace.cited(s.get("text") or "")}
    return {k: {"subject": f.get("subject"), "value": f.get("value"), "status": f.get("status")}
            for k, f in facts.items() if k in cited}


def docx(title: str, byline: str, written: list[tuple[rubric.Section, str]]) -> bytes:
    """Title, who it is for and when, a heading per part, and every [Add: …] gap highlighted."""
    return _document(title, byline, [(section.label, text) for section, text in written])


def form_docx(title: str, byline: str, filled: list[tuple[str, str]]) -> bytes:
    """The funder's form filled in: its own numbered headings, with what is left marked [Add: …] or [You fill: …]."""
    return _document(title, byline, filled)


def _document(title: str, byline: str, headed: list[tuple[str, str]]) -> bytes:
    from docx import Document
    from docx.enum.text import WD_COLOR_INDEX

    doc = Document()
    doc.add_heading(title, level=0)
    doc.add_paragraph().add_run(byline).italic = True
    for heading, text in headed:
        doc.add_heading(heading, level=1)
        for para in trace.paragraphs(text):
            for line in para.splitlines():
                if not line.strip():
                    continue
                bullet = bool(_BULLET.match(line))
                out = doc.add_paragraph(style="List Bullet" if bullet else None)
                line = _BULLET.sub("", line) if bullet else line.strip()
                for i, piece in enumerate(_LEFT.split(line)):
                    if piece:
                        run = out.add_run(piece)
                        if i % 2:
                            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
