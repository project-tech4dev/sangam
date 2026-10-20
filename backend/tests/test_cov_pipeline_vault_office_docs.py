"""Doc Vault reading and checks, Word/Excel text extraction, and the docs/ renderer, with no database."""

from __future__ import annotations

import contextlib
import io
from types import SimpleNamespace

import pytest
from docx import Document
from openpyxl import Workbook

from app.ingest import office, render_docs, vault
from app.integrations import storage


def docx_bytes(paras, table=None):
    d = Document()
    for p in paras:
        d.add_paragraph(p)
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, v in enumerate(row):
                t.cell(r, c).text = v
        t.cell(0, 0).merge(t.cell(0, 1))
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def xlsx_bytes(rows, title="Sheet"):
    wb = Workbook()
    ws = wb.active
    ws.title = title
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TestOffice:
    def test_docx_paragraphs_and_tables_in_order(self):
        data = docx_bytes(["Intro", "  ", "After"], table=[["A", "B", "C"], ["", "", ""], ["1", "2", "3"]])
        assert office.docx_paragraphs(data) == ["Intro", "After", "A\nB | C", "1 | 2 | 3"]  # a merged cell once

    def test_xlsx_markdown_caps_rows(self, monkeypatch):
        monkeypatch.setattr(office, "SHEET_ROWS", 2)
        data = xlsx_bytes([["a", None], [None, None], ["b", 2], ["c", 3]], title="Grants")
        assert office.xlsx_markdown(data) == "## Grants\na | \n… first 2 rows only"


DOC = {"id": 12, "title": "Annual report", "doc_type": "annual_report", "storage_path": "org-vault/3/ar.pdf"}


class TestVaultUrls:
    def test_url_round_trip(self):
        assert vault.document_id(vault.url(12)) == 12
        assert vault.document_id(vault.url(12) + "/") == 12
        assert vault.document_id(vault.URL_PREFIX + "abc") is None
        assert vault.document_id("https://example.org/12") is None
        assert vault.marked_ids("see [vault:3#p2] and [vault:4#para1]") == {3, 4}


class FakeConn:
    def __init__(self, rows):
        self.rows, self.sql = rows, []

    async def execute(self, sql, params=None):
        self.sql.append((sql, params))
        rows = self.rows

        class C:
            async def fetchall(self):
                return rows

            async def fetchone(self):
                return rows[0] if rows else None
        return C()


class FakePool:
    def __init__(self, rows):
        self.conn = FakeConn(rows)

    @contextlib.asynccontextmanager
    async def connection(self):
        yield self.conn


class TestVaultQueries:
    async def test_queries_are_scoped_to_the_vault_bucket(self):
        pool = FakePool([{"id": 1}, {"id": 2}])
        assert await vault.documents(pool, 3) == [{"id": 1}, {"id": 2}]
        assert await vault.document(pool, 1) == {"id": 1}
        assert await vault.in_vault(pool, 3, {2, 1}) == {1, 2}
        assert await vault.in_vault(pool, 3, set()) == set()
        assert all(p[-1] == "org-vault/%" for _, p in pool.conn.sql)
        assert pool.conn.sql[2][1][0] == [1, 2]

    async def test_document_rows(self):
        conn = FakeConn([{"id": 5}])
        out = await vault.document_rows(conn, [{"source_url": vault.url(5)}, {"source_url": "https://x.org"}])
        assert out == {vault.url(5): 5} and conn.sql[0][1][0] == [5]
        assert await vault.document_rows(conn, [{"source_url": "https://x.org"}]) == {}


class TestVaultRead:
    @pytest.fixture
    def download(self, monkeypatch):
        box = {}

        async def dl(bucket, key, max_bytes):
            return box.get("data")
        monkeypatch.setattr(storage, "download", dl)
        return box

    async def test_refusals(self, download):
        assert "no file on record" in await vault.read({**DOC, "storage_path": None})
        assert "is not a PDF, Word" in await vault.read({**DOC, "storage_path": "org-vault/3/x.txt"})
        assert "could not be downloaded" in await vault.read(DOC)

    async def test_pdf_pages_are_marked_and_cut(self, download, monkeypatch):
        download["data"] = b"%PDF"
        body = "--- page 1 ---\n" + ("line of text\n" * 3000) + "--- page 2 ---\nend"
        monkeypatch.setattr(vault, "extract_pdf_text", lambda path, s, e: (body, 2))
        out = await vault.read(DOC)
        assert "[vault:12#p1]" in out and "2 pages. Cite a fact" in out
        assert out.endswith("Call again with a later start_page.") and len(out) < vault.MAX_CHARS + 500

    async def test_a_scanned_pdf_says_so(self, download, monkeypatch):
        download["data"] = b"%PDF"
        monkeypatch.setattr(vault, "extract_pdf_text", lambda path, s, e: ("  ", 7))
        assert "7 pages and no text in that range: most likely scanned" in await vault.read(DOC)

    async def test_docx_paragraphs_are_paged(self, download, monkeypatch):
        download["data"] = docx_bytes([f"Para {i} " + "x" * 500 for i in range(80)])
        out = await vault.read({**DOC, "storage_path": "org-vault/3/plan.DOCX"}, start_page=3)
        assert out.splitlines()[3] == "" and "[vault:12#para3] Para 2" in out
        assert "… cut at paragraph" in out and "Call again with start_page=" in out

    async def test_xlsx_rows_are_cited_and_contacts_dropped(self, download):
        download["data"] = xlsx_bytes([["Grant", "Email"], ["h2", "x"], ["h3", "y"], ["Acme 5 lakh", "a@b.org"]])
        out = await vault.read({**DOC, "storage_path": "org-vault/3/g.xlsx"}, sheet="Sheet", limit=999)
        assert "[vault:12#Sheet!A4] A4='Acme 5 lakh'" in out and "a@b.org" not in out

    def test_xlsx_text_is_cut(self):
        data = xlsx_bytes([["h"], ["h"], ["h"]] + [[f"row {i} " + "z" * 250] for i in range(200)])
        out = vault.xlsx_text(DOC, data, "Sheet", 1, 200)
        assert out.endswith("Read one sheet.")


VURL = vault.url(4)


class TestVaultChecks:
    def test_check_findings(self):
        findings = [
            {"concept": "a", "source_url": "https://x.org"},
            {"concept": "b", "source_url": vault.url(9), "evidence": "[vault:9#p1]"},
            {"concept": "c", "source_url": VURL, "evidence": "no marker"},
            {"concept": "d", "source_url": VURL, "evidence": "[vault:4#p2]"},
        ]
        problems = vault.check_findings(findings, {4})
        assert [p.split(":")[0] for p in problems] == ["a", "b", "c"]

    def test_hold_at_promote(self):
        findings = [
            {"concept": "identity", "source_url": "https://x.org"},
            {"concept": "identity", "source_url": VURL, "evidence": "nothing"},
            {"concept": "contact", "source_url": VURL, "evidence": "[vault:4#p1]"},
            {"concept": "rfp", "source_url": VURL, "evidence": "[vault:4#p1]"},
        ]
        kept, held = vault.hold_at_promote(findings, "funder")
        assert kept == [findings[0], findings[3]]
        assert held == ["identity from document 4: no [vault:4#…] marker",
                        "contact from document 4: the organisation's own fact, kept private"]
        _, held_org = vault.hold_at_promote(findings[3:], "org")
        assert held_org == ["rfp from document 4: an organisation's own documents stay with it"]


class TestRenderDocs:
    def test_replace_block(self):
        text = "intro\n<!-- generated:x -->\nold\n<!-- /generated -->\noutro"
        assert render_docs.replace_block(text, "x", "new\n") == \
            "intro\n<!-- generated:x -->\nnew\n<!-- /generated -->\noutro"
        assert render_docs.replace_block("intro\n", "y", "b").endswith("<!-- generated:y -->\nb\n<!-- /generated -->\n")

    def test_client_block(self):
        assert "No organisation" in render_docs.client_block([])
        out = render_docs.client_block([{"canonical_name": None, "entity_name": "acme", "researched_on": "d",
                                         "status": "ok", "fact_count": 3, "slug": "acme"}])
        assert "| acme | d | ok | 3 | [entity file](../../entities/acme.md) |" in out

    async def test_main_needs_database_url(self, monkeypatch, capsys):
        monkeypatch.delenv("DATABASE_URL", raising=False)
        assert await render_docs.main() == 1
        assert "DATABASE_URL is not set" in capsys.readouterr().err

    async def test_main_writes_entities_and_client_blocks(self, monkeypatch, tmp_path, capsys):
        rows = [
            {"slug": "acme", "entity_kind": "funder", "status": "complete", "canonical_name": "Acme",
             "entity_name": "acme", "researched_on": "2026-01-01", "fact_count": 2},
            {"slug": "acme", "entity_kind": "funder", "status": "complete", "canonical_name": "Acme",
             "entity_name": "acme", "researched_on": "2026-02-01", "fact_count": 5},
            {"slug": "zeta", "entity_kind": "unknown", "status": "wrong_entity", "canonical_name": None,
             "entity_name": "zeta", "researched_on": "2026-01-01", "fact_count": 0},
        ]

        class Conn:
            async def __aenter__(self):
                return self

            async def __aexit__(self, *a):
                return False

            async def execute(self, sql, params=None):
                data = rows if "entity_research" in sql else [{"slug": "ngo-one", "name": "NGO One"}]

                class C:
                    async def fetchall(self):
                        return data
                return C()

        async def connect(url, row_factory):
            return Conn()
        built = []

        async def build(history):
            built.append(len(history))
            return f"# {history[-1]['slug']}\n"

        async def publish(slug, md):
            return slug if slug == "acme" else None
        clients = tmp_path / "clients"
        (clients / "ngo-one").mkdir(parents=True)
        (clients / "ngo-one" / "README.md").write_text("Hand written.\n")
        (clients / "draft").mkdir()
        (clients / "README.md").write_text("Index\n")
        monkeypatch.setenv("DATABASE_URL", "postgresql://x")
        monkeypatch.setattr(render_docs, "psycopg", SimpleNamespace(AsyncConnection=SimpleNamespace(connect=connect)))
        monkeypatch.setattr(render_docs, "build", build)
        monkeypatch.setattr(render_docs, "publish_markdown", publish)
        monkeypatch.setattr(render_docs, "ENTITIES", tmp_path / "entities")
        monkeypatch.setattr(render_docs, "CLIENTS", clients)

        assert await render_docs.main() == 0
        assert built == [2, 1]
        assert (tmp_path / "entities" / "acme.md").read_text() == "# acme\n"
        readme = (clients / "ngo-one" / "README.md").read_text()
        assert readme.startswith("Hand written.") and "| Acme | 2026-02-01 | complete | 5 |" in readme
        assert "zeta" not in readme
        assert "[NGO One](ngo-one/README.md) | 1 funders researched" in (clients / "README.md").read_text()
        out = capsys.readouterr().out
        assert "docs/entities: 2 file(s), 1 published" in out and "zeta (unknown/wrong_entity)" in out
        assert "refreshed 1 client file(s)" in out
