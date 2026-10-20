"""assemble, gates and records: the proposal as a document, step 1's requirements, and a registry fact's check."""

from __future__ import annotations

import io
from types import SimpleNamespace

import pytest
from docx import Document

from app.proposals import assemble, gates, records
from app.proposals.tools import record_tools
from app.repositories import profile_facts
from app.repositories import proposals as repo


def _work(**extra):
    return {"earlier_phase": False, "sections": {}, "facts": {}, **extra}


# assemble

def test_plain_drops_markers_and_bold():
    assert assemble.plain("**120** villages ⟦k1⟧ in Dewas ⟦k2⟧") == "120 villages in Dewas"
    assert assemble.plain(None) == ""


def test_parts_put_the_summary_first_and_skip_unwritten_parts():
    work = _work(sections={"A3": {"text": "Wells run dry ⟦k1⟧"}, "A1": {"text": "In short"}, "A2": {"text": ""}})
    assert [(s.code, text) for s, text in assemble.parts(work)] == [("A1", "In short"), ("A3", "Wells run dry")]


def test_final_uses_the_smoothed_wording_unless_the_part_changed_since():
    work = _work(sections={"A3": {"text": "raw three", "rev": 2}, "A4": {"text": "raw four", "rev": 5}},
                 polished={"parts": {"A3": "smooth three", "A4": "smooth four"}, "revs": {"A3": 2, "A4": 4}})
    assert [text for _, text in assemble.final(work)] == ["smooth three", "raw four"]
    assert [text for _, text in assemble.final(_work(sections={"A3": {"text": "x"}}))] == ["x"]


def test_snapshot_keeps_only_the_facts_the_text_cites():
    work = _work(sections={"A3": {"text": "120 villages ⟦k1⟧"}, "A4": {}},
                 facts={"k1": {"subject": "Villages", "value": "120", "status": "confirmed", "origin": "org"},
                        "k2": {"subject": "Budget", "value": "₹2 crore", "status": "proposed"}})
    assert assemble.snapshot(work) == {"k1": {"subject": "Villages", "value": "120", "status": "confirmed"}}


def test_docx_has_a_heading_per_part_bullets_and_highlighted_gaps():
    work = _work(sections={"A3": {"text": "Wells run dry.\n\n- 120 villages\n\n\n* [Add: survey year] and more"}})
    data = assemble.docx("Water for Dewas", "For Tata Trusts", assemble.parts(work))
    doc = Document(io.BytesIO(data))
    paras = [(p.style.name, p.text) for p in doc.paragraphs]
    assert paras[0] == ("Title", "Water for Dewas")
    assert ("Heading 1", "Problem and baseline") in paras
    assert ("List Bullet", "120 villages") in paras
    gap = next(p for p in doc.paragraphs if "survey year" in p.text)
    assert [r.font.highlight_color is not None for r in gap.runs] == [True, False]


def test_docx_skips_a_blank_line_handed_to_it_unstripped():
    section = assemble.parts(_work(sections={"A3": {"text": "x"}}))[0][0]
    doc = Document(io.BytesIO(assemble.docx("T", "B", [(section, "Wells run dry.\n   ")])))
    assert [p.text for p in doc.paragraphs][-1] == "Wells run dry."


# gates

def _detail(note="", not_judged=(), cites=(), relationship_points=None):
    return {"criteria": [{"code": "eligibility", "note": note, "cites": list(cites)},
                         {"code": "relationship", "points": relationship_points}],
            "not_judged": list(not_judged)}


def test_build_lists_csr1_only_for_section_135_money():
    csr = gates.build(_detail(), "corporate_csr")
    trust = gates.build(_detail(), "trust")
    bound = gates.build(_detail(), "trust", section_135_bound=True)
    assert [r["key"] for r in csr["requirements"]] == ["12a", "80g", "csr1", "audits"]
    assert [r["key"] for r in trust["requirements"]] == ["12a", "80g", "audits"]
    assert "csr1" in [r["key"] for r in bound["requirements"]]
    assert [r["hard"] for r in csr["requirements"]] == [False, False, True, False]


def test_build_says_what_the_match_found_on_record():
    detail = _detail(note="12A registered on 2015-03-01.; 80G not recorded", not_judged=["Audited accounts for 2023"],
                     cites=["[org:7]"])
    built = gates.build(detail, "trust")
    on_record = {r["key"]: r["on_record"] for r in built["requirements"]}
    assert on_record == {"12a": "12A registered on 2015-03-01", "80g": "80G not recorded",
                         "audits": "Audited accounts for 2023"}
    assert built["cites"] == ["[org:7]"]


def test_build_turns_the_funders_own_bar_into_a_requirement_clipped_for_the_screen():
    long = "x" * 300
    built = gates.build(_detail(not_judged=[f"{gates.BAR}: Three years in Madhya Pradesh", f"{gates.BAR}: {long}"]),
                        None)
    bars = [r for r in built["requirements"] if r["key"].startswith("bar")]
    assert [r["key"] for r in bars] == ["bar1", "bar2"]
    assert bars[0]["rule"] == "Three years in Madhya Pradesh"
    assert len(bars[1]["rule"]) == gates.CLIP and bars[1]["rule"].endswith("…")


def test_build_marks_an_earlier_phase_from_the_relationship_points():
    assert gates.build(_detail(relationship_points=10), None)["earlier_phase"] is True
    assert gates.build({"criteria": None}, None)["earlier_phase"] is False
    assert gates.build({}, None)["requirements"][0]["on_record"] == "Nothing on record."


# records

class _Pool:
    def __init__(self, rows):
        self.rows, self.calls = rows, 0

    def connection(self):
        pool = self

        class _Conn:
            async def __aenter__(self):
                pool.calls += 1
                return self

            async def __aexit__(self, *exc):
                return False

            async def execute(self, sql):
                return SimpleNamespace(fetchall=_fetch)

        async def _fetch():
            return pool.rows

        return _Conn()


async def test_districts_are_read_once_and_short_names_skipped(monkeypatch):
    monkeypatch.setattr(records, "_DISTRICTS", set())
    pool = _Pool([{"name": "dewas"}, {"name": "mon"}, {"name": "east-khasi  hills"}, {"name": None}])
    assert await records.districts(pool) == {"dewas", "east khasi hills"}
    await records.districts(pool)
    assert pool.calls == 1


def test_appears_compares_numbers_within_rounding_and_words_otherwise():
    assert records.appears("₹2.4 crore", ["Spent ₹2.40 crore in 2024"])
    assert not records.appears("₹2.5 crore", ["Spent ₹2.40 crore"])
    assert records.appears("Society", ["Jan Sahas · society · Dewas"])
    assert not records.appears("Trust", ["Jan Sahas · society"])


P = SimpleNamespace(org_id=7, org_name="Jan Sahas", funder_id=3, work={"match": {"id": 5}})


@pytest.fixture
def registry(monkeypatch):
    state = {"facts": object(), "match": {"id": 5}}

    async def org_facts(pool, caller_pool, org):
        return state["facts"]

    async def match_for(pool, org_id, funder_id):
        return state["match"]

    monkeypatch.setattr(profile_facts, "org_facts", org_facts)
    monkeypatch.setattr(repo, "match_for", match_for)
    monkeypatch.setattr(record_tools, "org_text", lambda facts: "Jan Sahas [org:7] · society\nFounded 2001 [org:7]")
    monkeypatch.setattr(record_tools, "match_text", lambda m: "Funder Match [score:5]: 62 of 100")
    return state


def _fact(marker, value):
    return {"subject": "Detail", "marker": marker, "value": value}


async def test_check_accepts_a_value_in_the_row_its_marker_names(registry):
    assert await records.check(None, P, _fact("[org:7]", "society")) is None
    assert await records.check(None, P, _fact("[score:5]", "62")) is None


@pytest.mark.parametrize("marker, value, why", [
    ("[org:7]", "trust", "'trust' is not what [org:7] says"),
    ("[org:8]", "x", "[org:8] is not Jan Sahas"),
    ("[score:6]", "62", "[score:6] is not the Funder Match this proposal started from"),
    ("[funder:3]", "x", "[funder:3] is not in Jan Sahas's record"),
])
async def test_check_refuses_with_the_reason(registry, marker, value, why):
    assert await records.check(None, P, _fact(marker, value)) == f"Detail: {why}"


async def test_check_refuses_when_the_row_is_gone(registry):
    registry["match"] = None
    assert await records.check(None, P, _fact("[score:5]", "62")) == "Detail: [score:5] is no longer held"
    registry["facts"] = None
    assert await records.check(None, P, _fact("[org:7]", "x")) == "Detail: Sangam holds no record for Jan Sahas"
