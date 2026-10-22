"""Filling the funder's form: the outline code keeps, the check that refuses an added claim, and every fallback."""

from __future__ import annotations

import base64
import io

import pytest
from docx import Document
from langchain_anthropic.chat_models import _format_messages
from langchain_core.messages import HumanMessage

from app.proposals import assemble, fill, form

CODES = {"A1", "A2", "A3", "A4", "A5", "A15"}


def heading(**kw):
    base = {"number": None, "label": "Heading", "ask": "", "kind": "prose", "limit": None, "uses": [], "only_if": None}
    return form.Heading(**{**base, **kw})


class TestTheOutline:
    def test_codes_follow_the_forms_order_and_unknown_parts_are_dropped(self):
        out = form.check(form.Outline(headings=[
            heading(number="1.", label="Executive Summary", uses=["A1"], limit=form.Limit(n=2000, unit="chars")),
            heading(number="5", label="Proposed strategy", uses=["A4", "A5", "A4", "Z9"]),
        ]), CODES)
        assert [(h["code"], h["number"], h["uses"]) for h in out] == [("T1", "1", ["A1"]), ("T2", "5", ["A4", "A5"])]
        assert out[0]["limit"] == {"n": 2000, "unit": "chars"} and out[1]["limit"] is None

    def test_only_prose_draws_on_parts_and_a_label_less_heading_joins_the_one_before(self):
        out = form.check(form.Outline(headings=[
            heading(label="Financial", ask="Outline the funding", kind="table", uses=["A15"]),
            heading(label="  ", ask="per stage of the project"),
        ]), CODES)
        assert len(out) == 1 and out[0]["uses"] == []
        assert out[0]["ask"] == "Outline the funding per stage of the project"

    def test_long_text_is_capped_and_brackets_cannot_close_a_mark(self):
        out = form.check(form.Outline(headings=[heading(label="x" * 500, ask="say [this] " * 300)]), CODES)
        assert len(out[0]["label"]) == form.LABEL and len(out[0]["ask"]) <= form.ASK
        assert "[" not in out[0]["ask"] and "]" not in out[0]["ask"]

    def test_a_pdf_reaches_anthropic_as_a_document(self):
        content = form.content("RF template", ".pdf", b"%PDF-1.4 fake")
        _, messages = _format_messages([HumanMessage(content)], model="claude-sonnet-5-5")
        block = messages[0]["content"][-1]
        assert block["type"] == "document"
        assert block["source"] == {"type": "base64", "media_type": "application/pdf",
                                   "data": base64.b64encode(b"%PDF-1.4 fake").decode()}

    def test_a_word_form_goes_as_its_text(self):
        doc = Document()
        doc.add_paragraph("1. Executive Summary (max 2000 characters)")
        buffer = io.BytesIO()
        doc.save(buffer)
        content = form.content("RF template", ".docx", buffer.getvalue())
        assert "1. Executive Summary (max 2000 characters)" in content[-1]["text"]

    def test_another_kind_of_file_is_refused(self):
        with pytest.raises(form.Unreadable, match="PDF or Word"):
            form.content("budget", ".xlsx", b"")


WRITTEN = {"A4": "Raise incomes of 2,400 women farmers in Goalpara.", "A5": "Run 54 trainings in 22 villages.",
           "A1": "SeSTA will work with 2,400 women."}
STRATEGY = {"code": "T5", "number": "5", "label": "Proposed strategy", "ask": "How the problem will be addressed",
            "kind": "prose", "limit": {"n": 90, "unit": "chars"}, "uses": ["A4", "A5"], "only_if": None}


ABOUT = "Organisation name: SeSTA\nCountry: India"
ALL = fill.everything(WRITTEN, ABOUT)


class TestTheFillCheck:
    def test_the_text_is_every_part_by_code_after_what_sangam_knows(self):
        assert ALL.startswith("Organisation name: SeSTA\nCountry: India\n\n[A4]\nRaise incomes")
        assert "[A1]\nSeSTA will work with 2,400 women." in ALL
        assert fill.everything({"A2": ""}) == ""

    def test_a_cut_or_a_figure_from_another_part_is_allowed(self):
        assert fill.check(ALL, "Run 54 trainings for 2,400 women.", STRATEGY["limit"], districts={"goalpara"}) == []

    def test_an_added_or_changed_figure_is_refused(self):
        assert fill.check(ALL, "Raise incomes of 2,500 women farmers.", STRATEGY["limit"]) == ["added 2500"]

    def test_an_answer_over_the_limit_is_refused(self):
        assert fill.check(ALL, "Raise incomes. " * 7, STRATEGY["limit"]) == ["ran to 104 characters; the form allows 90"]

    def test_a_new_gap_is_not_a_claim(self):
        assert fill.check("Run 54 trainings.", "Run 54 trainings. [Add: the trainers' names]", None) == []

    def test_the_limit_counts_plain_characters_or_words(self):
        assert fill.over({"n": 10, "unit": "chars"}, "1234567890 ⟦k1⟧") is None
        assert fill.over({"n": 10, "unit": "chars"}, "12345678901") == "11 characters; the form allows 10"
        assert fill.over({"n": 3, "unit": "words"}, "one two three four") == "4 words; the form allows 3"
        assert fill.over({"n": 1, "unit": "pages"}, "x" * 9000) is None
        assert fill.over(None, "x" * 9000) is None


NAME = {**STRATEGY, "code": "T2", "number": None, "label": "Organization Name", "ask": "", "kind": "field",
        "limit": None, "uses": []}
PAN = {**NAME, "code": "T9", "label": "PAN Card Number"}


def outline(*extra):
    return [{**STRATEGY, "limit": None, "code": "T1", "number": "1", "label": "Executive Summary", "uses": ["A1"]},
            STRATEGY, *extra]


class TestEveryHeadingComesOut:
    def resolve(self, headings, answers, earlier_phase=False):
        filled = fill.resolve(headings, WRITTEN, answers, earlier_phase=earlier_phase, about=ABOUT)
        return {f.heading["code"]: f for f in filled}

    def test_a_checked_answer_is_used_even_from_a_part_the_heading_was_not_matched_to(self):
        out = self.resolve(outline(), {"T5": "SeSTA will work with 2,400 women in 22 villages."})
        assert out["T5"].how == "answered" and out["T5"].text == "SeSTA will work with 2,400 women in 22 villages."

    def test_a_field_sangam_knows_is_filled_and_one_it_does_not_is_left(self):
        out = self.resolve([NAME, PAN], {"T2": "SeSTA", "T9": "[You fill: what the form asks]"})
        assert out["T2"].how == "answered" and out["T2"].text == "SeSTA"
        assert out["T9"].how == "you_fill" and out["T9"].text == "[You fill: PAN Card Number]"

    def test_an_answer_that_is_only_gaps_is_still_to_add(self):
        out = self.resolve(outline(), {"T5": "[Add: the indicator] [Add: the baseline].\n\n[Add: the products]."})
        assert out["T5"].how == "add" and out["T5"].text.startswith("[Add: the indicator]")

    def test_a_refused_answer_becomes_a_plain_mark_with_the_limit_and_is_logged(self, caplog):
        with caplog.at_level("INFO", logger="app.proposals.fill"):
            out = self.resolve(outline(), {"T5": "Raise incomes of 9,999 women."})
        assert out["T5"].how == "add"
        assert out["T5"].text == "[Add: How the problem will be addressed (the form allows 90 characters)]"
        assert "5. Proposed strategy refused: added 9999" in caplog.text

    def test_no_answer_is_a_mark_and_never_the_parts_pasted(self):
        out = self.resolve(outline(NAME), {})
        assert out["T1"].text == "[Add: How the problem will be addressed]"
        assert out["T2"].text == "[You fill: Organization Name]" and out["T2"].how == "you_fill"

    def test_tables_and_annexes_are_left_for_the_person_whatever_the_model_says(self):
        table = {**PAN, "code": "T10", "label": "Budget", "kind": "table"}
        assert self.resolve([table], {"T10": "SeSTA"})["T10"].text == "[You fill: Budget]"

    def test_skipped_headings_and_a_second_phase_question_on_a_first_phase_are_left_out(self):
        terms = {**STRATEGY, "code": "T10", "kind": "skip"}
        phase2 = {**STRATEGY, "code": "T11", "only_if": "earlier_phase", "uses": []}
        assert self.resolve([terms, phase2], {}) == {}
        assert set(self.resolve([terms, phase2], {}, earlier_phase=True)) == {"T11"}

    def test_the_report_names_headings_by_the_forms_numbers(self):
        filled = fill.resolve(outline(NAME, PAN), WRITTEN, {"T1": "SeSTA will work with 2,400 women.", "T2": "SeSTA"},
                              earlier_phase=False, about=ABOUT)
        assert fill.report(filled) == {"answered": ["1. Executive Summary", "Organization Name"],
                                       "add": ["5. Proposed strategy"], "you_fill": ["PAN Card Number"]}


def test_the_word_file_carries_the_forms_headings_and_marks_what_is_left():
    data = assemble.form_docx("Water for Dewas", "For Tata Trusts", [
        ("5. Proposed strategy", "Raise incomes. [Add: the trainers]"), ("PAN", "[You fill: PAN Card Number]")])
    doc = Document(io.BytesIO(data))
    assert [p.text for p in doc.paragraphs if p.style.name == "Heading 1"] == ["5. Proposed strategy", "PAN"]
    marked = [r.text for p in doc.paragraphs for r in p.runs if r.font.highlight_color]
    assert marked == ["[Add: the trainers]", "[You fill: PAN Card Number]"]


class FakeModel:
    """Answers a structured call with what it was given, and keeps the messages."""

    def __init__(self, result):
        self.result, self.calls = result, []

    def with_structured_output(self, schema):
        return self

    async def ainvoke(self, messages):
        self.calls.append(messages)
        if isinstance(self.result, Exception):
            raise self.result
        return self.result


def _pdf(pages):
    from pypdf import PdfWriter

    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=72, height=72)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


@pytest.fixture
def stored(monkeypatch):
    held = {"doc": {"id": 40, "title": "RF Proposal Template.PDF", "storage_path": "sangam-library/forms/rf.pdf"},
            "data": _pdf(2)}

    async def form_file(pool, document_id, org_id):
        held["asked"] = (document_id, org_id)
        return held["doc"]

    async def download(bucket, key, max_bytes):
        return held["data"]

    monkeypatch.setattr(form.repo, "form_file", form_file)
    monkeypatch.setattr(form.storage, "split", lambda path: tuple(path.split("/", 1)) if path else None)
    monkeypatch.setattr(form.storage, "download", download)
    return held


class TestReadingTheForm:
    async def test_the_outline_names_sangams_parts_and_keeps_the_forms_title(self, stored):
        from app.proposals import rubric

        model = FakeModel(form.Outline(headings=[heading(number="4", label="Problem Statement", uses=["A3"])]))
        sections = rubric.sections_for(False)
        out = await form.read(None, {"document_id": 40}, 7, sections, model)
        assert stored["asked"] == (40, 7)
        assert out["title"] == "RF Proposal Template" and out["document_id"] == 40 and out["pages_cut"] is False
        assert out["headings"][0]["uses"] == ["A3"]
        system = model.calls[0][0].content
        assert "{{PARTS}}" not in system and f"- A3 {sections[[s.code for s in sections].index('A3')].label}:" in system

    async def test_only_the_first_hundred_pages_are_sent(self, stored):
        stored["data"] = _pdf(form.MAX_PAGES + 1)
        model = FakeModel(form.Outline(headings=[heading(label="Problem")]))
        out = await form.read(None, {"document_id": 40}, 7, [], model)
        from pypdf import PdfReader

        sent = base64.b64decode(model.calls[0][1].content[-1]["base64"])
        assert out["pages_cut"] is True and len(PdfReader(io.BytesIO(sent)).pages) == form.MAX_PAGES

    @pytest.mark.parametrize("doc, data, outline, said", [
        (None, b"", [], "no copy"),
        ({"title": "rf.pdf", "storage_path": "sangam-library/rf.pdf"}, None, [], "could not be downloaded"),
        ({"title": "rf.pdf", "storage_path": "sangam-library/rf.pdf"}, _pdf(1), [heading(label="PAN", kind="field")],
         "no questions to answer in writing"),
    ])
    async def test_a_form_sangam_cannot_use_says_why(self, stored, doc, data, outline, said):
        stored.update(doc=doc, data=data)
        with pytest.raises(form.Unreadable, match=said):
            await form.read(None, {"document_id": 40}, 7, [], FakeModel(form.Outline(headings=outline)))


class TestTheFillCall:
    async def test_one_call_carries_all_the_text_and_every_heading_to_answer(self):
        model = FakeModel(fill.Answers(answers=[fill.Answer(code=" T5 ", text=" Raise incomes. "),
                                                fill.Answer(code="T1", text="  ")]))
        out = await fill.fill([STRATEGY, NAME, {**PAN, "code": "T10", "kind": "table"}], WRITTEN, "Plain words.",
                              model, ABOUT)
        assert out == {"T5": "Raise incomes."} and len(model.calls) == 1
        ask = model.calls[0][1].content
        assert ask.startswith("The style:\nPlain words.\n\n<text>\nOrganisation name: SeSTA")
        assert "[A5]\nRun 54 trainings in 22 villages." in ask
        assert ('<heading code="T5" kind="prose" title="5. Proposed strategy" limit="90 characters" parts="A4 A5">'
                "How the problem will be addressed</heading>") in ask
        assert 'code="T2" kind="field"' in ask and 'code="T10"' not in ask

    async def test_no_call_when_nothing_is_written_and_nothing_back_when_it_fails(self):
        idle = FakeModel(fill.Answers(answers=[]))
        assert await fill.fill([STRATEGY], {}, "", idle) == {} and idle.calls == []
        assert await fill.fill([STRATEGY], WRITTEN, "", FakeModel(RuntimeError("overloaded"))) == {}
